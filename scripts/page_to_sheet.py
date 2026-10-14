"""Reverse sync: copy what a live scouting page shows back into scouting-reports.xlsx.

Usage: python3 scripts/page_to_sheet.py sheet.xlsx slug [slug ...]

Use when the page is right and the sheet has drifted. Rewrites that player's
Summary/OFP/Follow/bio cells and every Grades, Traits and Data row from the
page's scout: block. Raw notes and other players are left alone.
"""
import re
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, PatternFill

INPUT = PatternFill("solid", fgColor="FFF6D6")
INPUT_COLS = {"Grades": {"Grade", "Sentences"}, "Traits": {"Form adjectives", "Sentences"}, "Data": set()}

ROOT = Path(__file__).resolve().parent.parent


def scout(slug):
    text = (ROOT / "scouting" / slug / "index.qmd").read_text()
    return yaml.safe_load(text.split("\n---", 1)[0].lstrip("-\n"))["scout"]


def header(ws):
    return [c.value for c in ws[1]]


def rewrite_rows(ws, slug, new_rows, raw_by_key, key_col):
    head = header(ws)
    raw_col = next((h for h in head if h and h.startswith("Raw")), None)
    keep = [[c.value for c in r] for r in ws.iter_rows(min_row=2) if r[0].value != slug]
    mine = []
    for rec in new_rows:
        if raw_col:
            rec[raw_col] = raw_by_key.get(rec.get(key_col), "")
        mine.append([rec.get(h, "") for h in head])
    # keep the player's rows where they were: insert at the first old position
    first = next((i for i, r in enumerate(ws.iter_rows(min_row=2)) if r[0].value == slug), len(keep))
    rows = keep[:first] + mine + keep[first:]
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append(r)
    # appended rows lose the sheet's look; put back wrap and the yellow input cells
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if head[c.column - 1] in INPUT_COLS.get(ws.title, set()):
                c.fill = INPUT


def main():
    path, slugs = sys.argv[1], sys.argv[2:]
    wb = openpyxl.load_workbook(path)
    for slug in slugs:
        sc = scout(slug)
        ws = wb["Players"]
        head = header(ws)
        for row in ws.iter_rows(min_row=2):
            if row[0].value != slug:
                continue
            for col, key in [("B/T", "bt"), ("Ht/Wt", "size"), ("Commit", "commit"), ("Travel team", "team"),
                             ("OFP", "ofp"), ("Follow", "follow"), ("Summary", "summary"),
                             ("PBR blog URL", "pbr_url")]:
                row[head.index(col)].value = str(sc.get(key) or "")
            row[head.index("Looks")].value = "; ".join(sc.get("looks") or [])

        def raw_map(sheet, key_col):
            h = header(wb[sheet])
            rc = next((x for x in h if x and x.startswith("Raw")), None)
            return {r[h.index(key_col)]: r[h.index(rc)] for r in wb[sheet].iter_rows(min_row=2, values_only=True)
                    if r[0] == slug} if rc else {}

        tools = [{"slug": slug, "Pitch / Tool": t["name"], "Grade": t.get("grade", ""),
                  "Form adjectives": "; ".join(t.get("tags") or []), "Velo": t.get("velo", ""),
                  "Spin": t.get("spin", ""), "Shape": t.get("shape", ""), "Sentences": t.get("notes", "")}
                 for t in sc.get("tools") or []]
        rewrite_rows(wb["Grades"], slug, tools, raw_map("Grades", "Pitch / Tool"), "Pitch / Tool")
        traits = [{"slug": slug, "Group": g["name"], "Form adjectives": "; ".join(g.get("tags") or []),
                   "Sentences": g.get("notes", "")} for g in sc.get("traits") or []]
        rewrite_rows(wb["Traits"], slug, traits, raw_map("Traits", "Group"), "Group")
        data = [{"slug": slug, "Measure": d["measure"], "Value": d["value"], "Source": d["source"]}
                for d in sc.get("data") or []]
        rewrite_rows(wb["Data"], slug, data, {}, "Measure")
        print(f"sheet updated from page: {slug}")
    wb.save(path)


if __name__ == "__main__":
    main()
